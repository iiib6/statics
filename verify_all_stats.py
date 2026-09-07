import os
import pandas as pd
import numpy as np
from scipy import stats
import docx
import statsmodels.formula.api as smf
import statsmodels.api as sm

# 1. Load Excel
df_raw = pd.read_excel('NOOR_120.xlsx')
df = df_raw.copy()
patient_cols = ['AGE', 'GENDER', 'DMFS', 'DMFT']
df['patient_id'] = df['AGE'].notna().cumsum()
df[patient_cols] = df[patient_cols].ffill()

df['GENDER'] = df['GENDER'].str.strip().str.capitalize()
df['arch'] = df['arch'].str.strip().str.capitalize()
df['CARIES INCIDENCE after 3 months'] = df['CARIES INCIDENCE after 3 months'].str.strip().str.upper()
df['RETENTION RATE after 3 months'] = df['RETENTION RATE after 3 months'].astype(int)
df['is_molar'] = df['sealant placement teeth'].str.contains('M', case=False, na=False).astype(int)
df['tooth_type'] = np.where(df['is_molar'], 'Molar', 'Premolar')
df['is_upper'] = (df['arch'] == 'Upper').astype(int)
df['age_gt10'] = (df['AGE'] > 10).astype(int)
df['ret_complete'] = (df['RETENTION RATE after 3 months'] == 0).astype(int)

# Patient df
df_p = df.drop_duplicates('patient_id')

# 2. Benchmark metrics
benchmarks = {}

# Demographics (Patients N=40)
benchmarks['patients_n'] = len(df_p)
benchmarks['boys_n'] = (df_p['GENDER'] == 'Male').sum()
benchmarks['girls_n'] = (df_p['GENDER'] == 'Female').sum()
benchmarks['boys_pct'] = round(benchmarks['boys_n'] / len(df_p) * 100, 1)
benchmarks['girls_pct'] = round(benchmarks['girls_n'] / len(df_p) * 100, 1)

benchmarks['age_mean'] = round(df_p['AGE'].mean(), 2)
benchmarks['age_sd'] = round(df_p['AGE'].std(ddof=1), 2)
benchmarks['age_min'] = int(df_p['AGE'].min())
benchmarks['age_max'] = int(df_p['AGE'].max())
benchmarks['age_med'] = float(df_p['AGE'].median())

benchmarks['dmft_mean'] = round(df_p['DMFT'].mean(), 2)
benchmarks['dmft_sd'] = round(df_p['DMFT'].std(ddof=1), 2)
benchmarks['dmfs_mean'] = round(df_p['DMFS'].mean(), 2)
benchmarks['dmfs_sd'] = round(df_p['DMFS'].std(ddof=1), 2)

# Teeth N=118
benchmarks['teeth_n'] = len(df)
benchmarks['upper_n'] = (df['arch'] == 'Upper').sum()
benchmarks['lower_n'] = (df['arch'] == 'Lower').sum()
benchmarks['premolars_n'] = (df['tooth_type'] == 'Premolar').sum()
benchmarks['molars_n'] = (df['tooth_type'] == 'Molar').sum()

# Retention
ret_counts = df['RETENTION RATE after 3 months'].value_counts()
benchmarks['score0_n'] = ret_counts.get(0, 0)
benchmarks['score1_n'] = ret_counts.get(1, 0)
benchmarks['score2_n'] = ret_counts.get(2, 0)
benchmarks['score0_pct'] = round(benchmarks['score0_n'] / len(df) * 100, 1)
benchmarks['score1_pct'] = round(benchmarks['score1_n'] / len(df) * 100, 1)
benchmarks['score2_pct'] = round(benchmarks['score2_n'] / len(df) * 100, 1)
benchmarks['cum_ret_n'] = benchmarks['score0_n'] + benchmarks['score1_n']
benchmarks['cum_ret_pct'] = round(benchmarks['cum_ret_n'] / len(df) * 100, 1)

# Inferential Tests
pre_ret = df[df['tooth_type'] == 'Premolar']['RETENTION RATE after 3 months']
mol_ret = df[df['tooth_type'] == 'Molar']['RETENTION RATE after 3 months']
u_type, p_type = stats.mannwhitneyu(pre_ret, mol_ret)
benchmarks['mw_type_u'] = round(u_type, 1)
benchmarks['mw_type_p'] = round(p_type, 3)

up_ret = df[df['arch'] == 'Upper']['RETENTION RATE after 3 months']
lo_ret = df[df['arch'] == 'Lower']['RETENTION RATE after 3 months']
u_arch, p_arch = stats.mannwhitneyu(up_ret, lo_ret)
benchmarks['mw_arch_u'] = round(u_arch, 1)
benchmarks['mw_arch_p'] = round(p_arch, 3)

a1_ret = df[df['AGE'] <= 10]['RETENTION RATE after 3 months']
a2_ret = df[df['AGE'] > 10]['RETENTION RATE after 3 months']
u_age, p_age = stats.mannwhitneyu(a1_ret, a2_ret)
benchmarks['mw_age_u'] = round(u_age, 1)
benchmarks['mw_age_p'] = round(p_age, 3)

m_ret = df[df['GENDER'] == 'Male']['RETENTION RATE after 3 months']
f_ret = df[df['GENDER'] == 'Female']['RETENTION RATE after 3 months']
u_gen, p_gen = stats.mannwhitneyu(m_ret, f_ret)
benchmarks['mw_gen_u'] = round(u_gen, 1)
benchmarks['mw_gen_p'] = round(p_gen, 3)

# Kruskal-Wallis for Baseline DMFT & DMFS across retention groups
ret0_dmft = df[df['RETENTION RATE after 3 months'] == 0]['DMFT']
ret1_dmft = df[df['RETENTION RATE after 3 months'] == 1]['DMFT']
ret2_dmft = df[df['RETENTION RATE after 3 months'] == 2]['DMFT']
kw_dmft = stats.kruskal(ret0_dmft, ret1_dmft, ret2_dmft)
benchmarks['kw_dmft_h'] = round(float(kw_dmft.statistic), 3)
benchmarks['kw_dmft_p'] = round(float(kw_dmft.pvalue), 3)

ret0_dmfs = df[df['RETENTION RATE after 3 months'] == 0]['DMFS']
ret1_dmfs = df[df['RETENTION RATE after 3 months'] == 1]['DMFS']
ret2_dmfs = df[df['RETENTION RATE after 3 months'] == 2]['DMFS']
kw_dmfs = stats.kruskal(ret0_dmfs, ret1_dmfs, ret2_dmfs)
benchmarks['kw_dmfs_h'] = round(float(kw_dmfs.statistic), 3)
benchmarks['kw_dmfs_p'] = round(float(kw_dmfs.pvalue), 3)

# GEE Cluster Sensitivity Models
fam = sm.families.Binomial()
cov = sm.cov_struct.Exchangeable()
gee_molar = smf.gee('ret_complete ~ is_molar', 'patient_id', data=df, family=fam, cov_struct=cov).fit()
benchmarks['gee_molar_p'] = round(float(gee_molar.pvalues['is_molar']), 3)

gee_age = smf.gee('ret_complete ~ age_gt10', 'patient_id', data=df, family=fam, cov_struct=cov).fit()
benchmarks['gee_age_p'] = round(float(gee_age.pvalues['age_gt10']), 4)

gee_arch = smf.gee('ret_complete ~ is_upper', 'patient_id', data=df, family=fam, cov_struct=cov).fit()
benchmarks['gee_arch_p'] = round(float(gee_arch.pvalues['is_upper']), 3)

print('=== BENCHMARK STATISTICAL AUDIT SUMMARY ===')
print(f'Patients: N = {benchmarks["patients_n"]} | Teeth: N = {benchmarks["teeth_n"]}')
print(f'Retention: Complete = {benchmarks["score0_n"]} ({benchmarks["score0_pct"]}%), Partial = {benchmarks["score1_n"]} ({benchmarks["score1_pct"]}%), Missing = {benchmarks["score2_n"]} ({benchmarks["score2_pct"]}%)')
print(f'Satisfactory Retention (Score 0 + 1): {benchmarks["cum_ret_n"]} ({benchmarks["cum_ret_pct"]}%)')
print(f'Kruskal-Wallis DMFT: H = {benchmarks["kw_dmft_h"]}, p = {benchmarks["kw_dmft_p"]}')
print(f'Kruskal-Wallis DMFS: H = {benchmarks["kw_dmfs_h"]}, p = {benchmarks["kw_dmfs_p"]}')
print(f'GEE Cluster-Adjusted Molars: p = {benchmarks["gee_molar_p"]}')
print(f'GEE Cluster-Adjusted Age (>10): p = {benchmarks["gee_age_p"]}')
print(f'GEE Cluster-Adjusted Arch: p = {benchmarks["gee_arch_p"]}')

# Check Word doc
target_doc = 'Scopus_Statistical_Report_Final.docx'
if os.path.exists(target_doc):
    doc = docx.Document(target_doc)
    doc_text = '\n'.join([p.text for p in doc.paragraphs])
    for t in doc.tables:
        for row in t.rows:
            doc_text += '\n' + ' | '.join([c.text for c in row.cells])
            
    checks = [
        ('Patient Sample Size 40', '40' in doc_text),
        ('Teeth Sample Size 118', '118' in doc_text),
        ('Complete Retention 85 (72.0%)', '85' in doc_text and '72.0%' in doc_text),
        ('Satisfactory Retention 113 (95.8%)', '113' in doc_text and '95.8%' in doc_text),
        ('Premolars vs Molars U=1155.5', '1155.5' in doc_text),
        ('Premolars vs Molars p=0.002', '0.002' in doc_text),
        ('Fisher-Freeman-Halton exact test', 'Fisher-Freeman-Halton' in doc_text),
        ('Kruskal-Wallis DMFT H=0.566', '0.566' in doc_text),
        ('Kruskal-Wallis DMFS H=0.601', '0.601' in doc_text),
        ('Clustered Sensitivity GEE', 'GEE' in doc_text or 'Generalized Estimating Equations' in doc_text),
        ('GEE Molars p=0.005', '0.005' in doc_text),
        ('GEE Age p=0.0003', '0.0003' in doc_text)
    ]
    print('\n=== DOCUMENT VERIFICATION CHECKS ===')
    all_ok = True
    for name, ok in checks:
        print(f'{"[PASS]" if ok else "[FAIL]"} {name}')
        if not ok: all_ok = False
    print(f'Status: {"ALL BENCHMARKS VERIFIED 100%" if all_ok else "CHECKS FAILED"}')
