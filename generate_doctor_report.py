import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_doctor_report():
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        p_footer = footer.paragraphs[0]
        p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_f = p_footer.add_run("Statistical Analysis & Clinical Findings Report | Scopus Manuscript Data | Page ")
        run_f.font.name = "Calibri"
        run_f.font.size = Pt(9)
        run_f.font.color.rgb = RGBColor(120, 120, 120)

    # Styles
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(40, 40, 40)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    def set_cell_border(cell, **kwargs):
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge_data = kwargs.get(edge)
            if edge_data:
                tag = f'w:{edge}'
                element = parse_xml(f'<{tag} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", 4)}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>')
                tcBorders.append(element)
            else:
                tag = f'w:{edge}'
                element = parse_xml(f'<{tag} {nsdecls("w")} w:val="none"/>')
                tcBorders.append(element)
        tcPr.append(tcBorders)

    def set_cell_shading(cell, color_hex):
        shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
        cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

    def add_custom_heading(text, level):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        if level == 1:
            run.font.size = Pt(15)
            run.font.color.rgb = RGBColor(20, 60, 120)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
        elif level == 2:
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(40, 90, 150)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        elif level == 3:
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(60, 60, 60)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
        return p

    def style_academic_table(table, col_widths, alignments):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            if i == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                
            for j, cell in enumerate(row.cells):
                cell.width = Inches(col_widths[j])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                tcPr = cell._tc.get_or_add_tcPr()
                tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="120" w:type="dxa"/><w:bottom w:w="120" w:type="dxa"/><w:left w:w="150" w:type="dxa"/><w:right w:w="150" w:type="dxa"/></w:tcMar>')
                tcPr.append(tcMar)
                
                if i == 0:
                    set_cell_border(cell, top={'sz': 12, 'val': 'single', 'color': '143C78'},
                                          bottom={'sz': 10, 'val': 'single', 'color': '143C78'})
                    set_cell_shading(cell, 'F0F4F8')
                elif i == len(table.rows) - 1:
                    set_cell_border(cell, bottom={'sz': 12, 'val': 'single', 'color': '143C78'})
                else:
                    set_cell_border(cell, bottom={'sz': 4, 'val': 'single', 'color': 'EAEAEA'})
                
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.05
                    p.alignment = alignments[j]
                    for r in p.runs:
                        r.font.name = 'Calibri'
                        r.font.size = Pt(9.5)
                        if i == 0:
                            r.bold = True
                            r.font.color.rgb = RGBColor(20, 60, 120)

    # -------------------------------------------------------------
    # DOCUMENT HEADER & METADATA
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(3)
    run_t = p_title.add_run("STATISTICAL RESULTS & CLINICAL FINDINGS REPORT")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(17)
    run_t.bold = True
    run_t.font.color.rgb = RGBColor(20, 60, 120)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    run_sub = p_sub.add_run("Clinical Evaluation of Pit and Fissure Sealant Retention and Associated Caries Incidence: A 3-Month Longitudinal Investigation\nPrepared for International Journal Publication (Scopus Indexed)")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(70, 70, 70)

    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in meta_table.rows:
        for c in r.cells:
            set_cell_shading(c, "F8F9FA")
            set_cell_border(c, top={'sz': 4, 'val': 'single', 'color': 'D0D5DD'},
                               bottom={'sz': 4, 'val': 'single', 'color': 'D0D5DD'},
                               left={'sz': 4, 'val': 'single', 'color': 'D0D5DD'},
                               right={'sz': 4, 'val': 'single', 'color': 'D0D5DD'})
    meta_table.rows[0].cells[0].paragraphs[0].add_run("Study Cohort: ").bold = True
    meta_table.rows[0].cells[0].paragraphs[0].add_run("40 Patients (29 Boys, 11 Girls; Ages 8–12)")
    meta_table.rows[0].cells[1].paragraphs[0].add_run("Evaluated Units: ").bold = True
    meta_table.rows[0].cells[1].paragraphs[0].add_run("118 Sound Permanent Teeth (77 Premolars, 41 Molars)")
    
    meta_table.rows[1].cells[0].paragraphs[0].add_run("Baseline Caries: ").bold = True
    meta_table.rows[1].cells[0].paragraphs[0].add_run("Mean DMFT = 2.58 ± 1.24 | Mean DMFS = 4.88 ± 2.21")
    meta_table.rows[1].cells[1].paragraphs[0].add_run("Follow-Up Period: ").bold = True
    meta_table.rows[1].cells[1].paragraphs[0].add_run("3 Months Post-Application")

    meta_table.rows[2].cells[0].paragraphs[0].add_run("Evaluation Index: ").bold = True
    meta_table.rows[2].cells[0].paragraphs[0].add_run("Simonsen's Criteria (Score 0: Complete, 1: Partial, 2: Lost)")
    meta_table.rows[2].cells[1].paragraphs[0].add_run("Statistical Framework: ").bold = True
    meta_table.rows[2].cells[1].paragraphs[0].add_run("Tooth-Level (Mann-Whitney U, Kruskal-Wallis, Fisher-Freeman-Halton) & Clustered Sensitivity (GEE)")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 1. SAMPLE SUMMARY & BASELINE CHARACTERISTICS
    # -------------------------------------------------------------
    add_custom_heading("1. Baseline Sample Characteristics & Anatomical Distribution", level=1)
    
    p = doc.add_paragraph("A total of 40 pediatric patients (29 boys [72.5%] and 11 girls [27.5%]) aged 8 to 12 years (mean age: 10.12 ± 1.24 years; median: 10.0 years) participated in this 3-month longitudinal study. Baseline epidemiological assessment demonstrated a mean DMFT index of 2.58 ± 1.24 (median: 3.0; range: 0–5) and a mean DMFS index of 4.88 ± 2.21 (median: 5.0; range: 0–8). Among the participants, 95.0% (n = 38) were caries-active at baseline, reflecting a clinically relevant pediatric cohort with genuine occlusal caries susceptibility.")
    
    p = doc.add_paragraph("A total of 118 sound permanent teeth received light-cured resin-based fissure sealants (mean: 2.95 ± 1.80 teeth per child; range: 1–8 teeth). The anatomical distribution comprised 71 maxillary teeth (60.2%) and 47 mandibular teeth (39.8%). By tooth class, premolars accounted for 65.3% (n = 77; including 21 RFP, 19 LFP, 17 RSP, and 20 LSP), while permanent molars accounted for 34.7% (n = 41; including 22 RFM, 18 LFM, and 1 RSM).")

    # Table 1: Baseline Demographics & Tooth Profile
    add_custom_heading("Table 1. Demographic profile of participants and anatomical distribution of sealed permanent teeth", level=3)
    t1_headers = ["Variable / Classification", "Subcategory / Tooth Code", "N (%)", "Summary Metric (Mean ± SD, Range)"]
    t1_data = [
        ["Patient Gender (N = 40)", "Male (Boys)\nFemale (Girls)", "29 (72.5%)\n11 (27.5%)", "Ratio = 2.64 : 1"],
        ["Patient Age (N = 40)", "≤ 10 Years (8–10)\n> 10 Years (11–12)", "24 (60.0%)\n16 (40.0%)", "Mean = 10.12 ± 1.24 years\nRange = 8.0 – 12.0 years"],
        ["Baseline Caries (N = 40)", "Caries-Active (DMFT > 0)\nCaries-Free (DMFT = 0)", "38 (95.0%)\n2 (5.0%)", "DMFT = 2.58 ± 1.24 (Range: 0–5)\nDMFS = 4.88 ± 2.21 (Range: 0–8)"],
        ["Jaw Location (N = 118)", "Maxillary (Upper Arch)\nMandibular (Lower Arch)", "71 (60.2%)\n47 (39.8%)", "—"],
        ["Tooth Type (N = 118)", "Premolars (Permanent)\nMolars (Permanent)", "77 (65.3%)\n41 (34.7%)", "Teeth per child = 2.95 ± 1.80\n(Range: 1 – 8)"],
        ["Specific Molar Positions", "Right First Molar (RFM)\nLeft First Molar (LFM)\nRight Second Molar (RSM)", "22 (18.6%)\n18 (15.3%)\n1 (0.8%)", "Total Molars = 41 (34.7%)"],
        ["Specific Premolar Positions", "Right First Premolar (RFP)\nLeft Second Premolar (LSP)\nLeft First Premolar (LFP)\nRight Second Premolar (RSP)", "21 (17.8%)\n20 (16.9%)\n19 (16.1%)\n17 (14.4%)", "Total Premolars = 77 (65.3%)"]
    ]
    t1 = doc.add_table(rows=len(t1_data) + 1, cols=4)
    for j, h in enumerate(t1_headers):
        t1.rows[0].cells[j].paragraphs[0].text = h
    row_idx = 1
    for item in t1_data:
        t1.rows[row_idx].cells[0].paragraphs[0].text = item[0]
        t1.rows[row_idx].cells[1].paragraphs[0].text = item[1]
        t1.rows[row_idx].cells[2].paragraphs[0].text = item[2]
        t1.rows[row_idx].cells[3].paragraphs[0].text = item[3]
        row_idx += 1
    style_academic_table(t1, [2.0, 2.2, 1.2, 1.8], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 2. OVERALL CLINICAL OUTCOMES
    # -------------------------------------------------------------
    add_custom_heading("2. Primary Clinical Outcomes at 3-Month Follow-Up", level=1)
    
    p = doc.add_paragraph("At the 3-month post-placement clinical evaluation, retention outcomes were categorized according to Simonsen's criteria into distinct levels:")
    
    bp_r1 = doc.add_paragraph(style='List Bullet')
    bp_r1.add_run("Complete Retention (Score 0): ").bold = True
    bp_r1.add_run("85 teeth (72.0%) showed complete coverage with all pit and fissure grooves fully sealed.")
    
    bp_r2 = doc.add_paragraph(style='List Bullet')
    bp_r2.add_run("Partial Retention (Score 1): ").bold = True
    bp_r2.add_run("28 teeth (23.7%) exhibited partial sealant loss with portions of the grooves exposed, while the remaining restoration remained bonded.")

    bp_r3 = doc.add_paragraph(style='List Bullet')
    bp_r3.add_run("Complete Loss (Score 2): ").bold = True
    bp_r3.add_run("5 teeth (4.2%) suffered total sealant dislodgement with complete loss of the protective resin material.")

    p_cum = doc.add_paragraph("Defining clinical success comprehensively, satisfactory retention (Score 0 + Score 1) was observed in 95.8% of teeth (n = 113). Concurrently, 95.8% (n = 113) of evaluated teeth remained caries-free (sound), whereas new occlusal carious lesions developed in 4.2% (n = 5) of the evaluated teeth.")

    # Table 2: Overall Condition
    add_custom_heading("Table 2. Overall retention distribution (Simonsen's criteria) and caries incidence at 3 months (N = 118)", level=3)
    t2 = doc.add_table(rows=7, cols=3)
    t2_headers = ["Clinical Parameter", "Simonsen Criteria / Clinical Category", "N (%)"]
    for j, h in enumerate(t2_headers):
        t2.rows[0].cells[j].paragraphs[0].text = h
    t2_data = [
        ["Sealant Retention Status", "Score 0: Complete Retention", "85 (72.0%)"],
        ["", "Score 1: Partial Retention", "28 (23.7%)"],
        ["", "Score 2: Complete Loss", "5 (4.2%)"],
        ["Satisfactory Retention", "Combined Score 0 + Score 1", "113 (95.8%)"],
        ["3-Month Caries Status", "Caries-Free (Sound Tooth)", "113 (95.8%)"],
        ["", "Caries Incidence (Decayed)", "5 (4.2%)"]
    ]
    for i, row_data in enumerate(t2_data):
        for j, val in enumerate(row_data):
            t2.rows[i+1].cells[j].paragraphs[0].text = val
    style_academic_table(t2, [2.5, 3.2, 1.5], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])
    
    p_t2_note = doc.add_paragraph("Note: Complete retention alone was 72.0% (n = 85); when including partially retained sealants that still offered clinical coverage, satisfactory retention was 95.8% (n = 113).")
    p_t2_note.runs[0].font.size = Pt(8.5)
    p_t2_note.runs[0].font.italic = True
    p_t2_note.paragraph_format.space_after = Pt(8)

    # -------------------------------------------------------------
    # 3. STATISTICAL FINDINGS & CLINICAL INTERPRETATIONS
    # -------------------------------------------------------------
    add_custom_heading("3. Statistical Analysis of Findings & Clinical Deductions", level=1)

    # FINDING 1: Tooth Type
    add_custom_heading("3.1 Finding 1: Anatomical Determinants — Premolars vs. Molars", level=2)
    p = doc.add_paragraph("Statistical Result: A highly statistically significant difference in sealant retention was detected between permanent premolars and permanent molars (Mann-Whitney U = 1155.5, Z = -3.054, p = 0.002; Pearson Chi-Square χ² = 13.291, df = 2, p = 0.0013).")
    
    p = doc.add_paragraph("Observed Data: Premolars achieved an 80.5% complete retention rate (n = 62/77) and experienced zero complete loss (0.0%). In sharp contrast, molars exhibited a 56.1% complete retention rate (n = 23/41), a 31.7% partial loss rate (n = 13/41), and accounted for 100% of all complete loss cases (12.2%, n = 5/41).")
    
    p = doc.add_paragraph("Clinical & Biological Rationale for the Manuscript:")
    bp_t1 = doc.add_paragraph(style='List Bullet')
    bp_t1.add_run("1. Fissure Morphology & Depth: ").bold = True
    bp_t1.add_run("Permanent molars possess complex, constricted, and deep 'I-type' and 'bottle-shaped' fissures that retain organic remnants and prevent deep etchant and resin penetration. Conversely, premolars have shallower, broader, V-shaped grooves that allow comprehensive acid etching and complete resin flow.")
    
    bp_t2 = doc.add_paragraph(style='List Bullet')
    bp_t2.add_run("2. Moisture Control & Saliva Contamination: ").bold = True
    bp_t2.add_run("Isolation in the posterior molar region is technically demanding in pediatric patients due to pooling of saliva in the posterior vestibule and active tongue movements, increasing the risk of micro-salivary contamination during the sensitive bonding procedure.")

    bp_t3 = doc.add_paragraph(style='List Bullet')
    bp_t3.add_run("3. Eruption Stage & Operculum: ").bold = True
    bp_t3.add_run("Newly emerged first permanent molars frequently present partially covered distal marginal ridges by a gingival operculum, compromising complete peripheral seal placement and mechanical retention.")

    bp_t4 = doc.add_paragraph(style='List Bullet')
    bp_t4.add_run("4. Masticatory Load Concentration: ").bold = True
    bp_t4.add_run("Permanent first molars absorb the highest vertical and lateral chewing forces, subjecting molar sealants to greater shearing stress than premolars.")

    # Table 3: Tooth Type & Jaw Location
    add_custom_heading("Table 3. Comparison of sealant retention across Anatomical Factors: Tooth Type and Jaw Location (N = 118)", level=3)
    t3 = doc.add_table(rows=5, cols=8)
    t3_headers = ["Anatomical Factor", "N Evaluated", "Score 0 (Complete)\nN (%)", "Score 1 (Partial)\nN (%)", "Score 2 (Missing)\nN (%)", "Mean Rank", "Mann-Whitney U (Z)", "p-value"]
    for j, h in enumerate(t3_headers):
        t3.rows[0].cells[j].paragraphs[0].text = h
    t3_data = [
        ["Premolars (Permanent)", "77", "62 (80.5%)", "15 (19.5%)", "0 (0.0%)", "54.01", "U = 1155.5\nZ = -3.054", "p = 0.002 **\n(Significant)"],
        ["Molars (Permanent)", "41", "23 (56.1%)", "13 (31.7%)", "5 (12.2%)", "69.82", "", ""],
        ["Maxillary (Upper Arch)", "71", "52 (73.2%)", "17 (23.9%)", "2 (2.8%)", "58.58", "U = 1603.5\nZ = -0.456", "p = 0.648\n(NS)"],
        ["Mandibular (Lower Arch)", "47", "33 (70.2%)", "11 (23.4%)", "3 (6.4%)", "60.88", "", ""]
    ]
    for i, row_data in enumerate(t3_data):
        for j, val in enumerate(row_data):
            t3.rows[i+1].cells[j].paragraphs[0].text = val
    style_academic_table(t3, [1.6, 0.8, 1.0, 1.0, 1.0, 0.7, 1.1, 0.9], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Insert Figure 1
    if os.path.exists('figure1_retention_tooth_type.png'):
        p_fig1 = doc.add_paragraph()
        p_fig1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture('figure1_retention_tooth_type.png', width=Inches(5.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1 = doc.add_paragraph()
        p_cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap1.paragraph_format.space_after = Pt(8)
        r_cap1 = p_cap1.add_run("Figure 1. Retention distribution of sealants in premolars and permanent molars after 3 months (p = 0.002).")
        r_cap1.bold = True
        r_cap1.font.size = Pt(9.5)
        r_cap1.font.name = "Calibri"

    # FINDING 2: Jaw Location
    add_custom_heading("3.2 Finding 2: Invariance Across Jaw Location (Maxillary vs. Mandibular Arches)", level=2)
    p = doc.add_paragraph("Statistical Result: No statistically significant difference in sealant retention was observed between maxillary and mandibular arches (Mann-Whitney U = 1603.5, Z = -0.456, p = 0.648; Pearson Chi-Square χ² = 0.888, df = 2, p = 0.641).")
    
    p = doc.add_paragraph("Observed Data: Maxillary teeth exhibited 73.2% complete retention (n = 52/71) and 2.8% complete loss (n = 2/71), while mandibular teeth exhibited 70.2% complete retention (n = 33/47) and 6.4% complete loss (n = 3/47). Mean ranks were closely matched (58.58 for upper vs. 60.88 for lower).")
    
    p = doc.add_paragraph("Clinical Interpretation: These findings demonstrate that standardized quadrant isolation with cotton rolls and saliva ejectors provides equivalent moisture control in both arches, confirming that arch location does not compromise clinical retention when standardized clinical isolation protocols are rigorously maintained.")

    # Insert Figure 2
    if os.path.exists('figure2_retention_jaw_location.png'):
        p_fig2 = doc.add_paragraph()
        p_fig2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture('figure2_retention_jaw_location.png', width=Inches(5.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2 = doc.add_paragraph()
        p_cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap2.paragraph_format.space_after = Pt(8)
        r_cap2 = p_cap2.add_run("Figure 2. Distribution of sealant retention according to jaw location after 3 months (p = 0.648).")
        r_cap2.bold = True
        r_cap2.font.size = Pt(9.5)
        r_cap2.font.name = "Calibri"

    # FINDING 3: Age & Gender
    add_custom_heading("3.3 Finding 3: Demographic Influences — Age Significance and Gender Equivalence", level=2)
    p = doc.add_paragraph("Statistical Result for Age: Sealant retention differed significantly between age groups (Mann-Whitney U = 1328.0, Z = -2.726, p = 0.006). Older children (> 10 years; 11–12 years) achieved 81.5% complete retention (n = 53/65) and 0.0% complete loss, whereas younger children (≤ 10 years; 8–10 years) achieved 60.4% complete retention (n = 32/53) and accounted for all 5 cases of complete sealant loss (9.4%).")
    
    p = doc.add_paragraph("Clinical & Behavioral Rationale for Age: Older children exhibit superior chairside behavioral cooperation, facilitating uninterrupted etching, washing, and drying. Furthermore, premolars and second molars in 11–12-year-olds have reached full clinical crown emergence, avoiding subgingival opercula and moisture pooling commonly encountered in recently erupted molars of younger children.")

    p = doc.add_paragraph("Statistical Result for Gender: Sealant retention showed no statistically significant difference between male and female patients (Mann-Whitney U = 1441.5, Z = -1.221, p = 0.222). Complete retention was 76.0% in boys (n = 57/75) and 65.1% in girls (n = 28/43), confirming clinical and biological equivalence.")

    # Table 4: Demographic Factors
    add_custom_heading("Table 4. Comparison of sealant retention across Demographic Subgroups (N = 118)", level=3)
    t4 = doc.add_table(rows=5, cols=8)
    t4_headers = ["Demographic Factor", "N Evaluated", "Score 0 (Complete)\nN (%)", "Score 1 (Partial)\nN (%)", "Score 2 (Missing)\nN (%)", "Mean Rank", "Mann-Whitney U (Z)", "p-value"]
    for j, h in enumerate(t4_headers):
        t4.rows[0].cells[j].paragraphs[0].text = h
    t4_data = [
        ["Age ≤ 10 Years (8–10)", "53", "32 (60.4%)", "16 (30.2%)", "5 (9.4%)", "66.94", "U = 1328.0\nZ = -2.726", "p = 0.006 **\n(Significant)"],
        ["Age > 10 Years (11–12)", "65", "53 (81.5%)", "12 (18.5%)", "0 (0.0%)", "53.43", "", ""],
        ["Male (Boys)", "75", "57 (76.0%)", "15 (20.0%)", "3 (4.0%)", "57.22", "U = 1441.5\nZ = -1.221", "p = 0.222\n(NS)"],
        ["Female (Girls)", "43", "28 (65.1%)", "13 (30.2%)", "2 (4.7%)", "63.48", "", ""]
    ]
    for i, row_data in enumerate(t4_data):
        for j, val in enumerate(row_data):
            t4.rows[i+1].cells[j].paragraphs[0].text = val
    style_academic_table(t4, [1.6, 0.8, 1.0, 1.0, 1.0, 0.7, 1.1, 0.9], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # FINDING 4: Retention vs Caries
    add_custom_heading("3.4 Finding 4: Association Between Sealant Retention Status and 3-Month Caries Incidence", level=2)
    p = doc.add_paragraph("Statistical Result: Cross-tabulation revealed a statistically significant association between sealant retention status and 3-month caries incidence. Because several contingency cells contained zero counts and expected frequencies below 5, the Fisher-Freeman-Halton Exact Test was established as the primary statistical test (Exact p < 0.0001, exact probability = 5.75 × 10⁻⁹). Pearson's Chi-Square test was computed as secondary supporting evidence (χ² = 118.00, df = 2, asymptotic p < 0.0001).")
    
    p = doc.add_paragraph("Observed Data: Teeth with completely retained sealants (Score 0, n = 85) and partially retained sealants (Score 1, n = 28) demonstrated a 100% caries-free rate (0.0% caries) during the 3-month follow-up period. Conversely, all 5 newly detected carious lesions occurred exclusively in teeth that experienced complete sealant loss (Score 2, n = 5).")
    
    p = doc.add_paragraph("Clinical Interpretation & Observational Association: These findings demonstrate that sealant retention status was significantly associated with caries incidence at 3 months. Intact and partially retained resin sealants provided a clinical physical barrier over vulnerable fissures, whereas complete dislodgement re-exposed previously conditioned enamel to oral cariogenic factors. However, because this was a 3-month observational study without an untreated control group, these findings establish a strong clinical association rather than absolute causal prevention, as oral hygiene habits, dietary patterns, and salivary factors also contribute to caries development over time.")

    # Table 5: Retention vs Caries
    add_custom_heading("Table 5. Contingency table of Sealant Retention Status vs. 3-Month Caries Incidence (N = 118)", level=3)
    t5 = doc.add_table(rows=5, cols=6)
    t5_headers = ["Retention Status (Simonsen's)", "Caries-Free (Sound)\nN (%)", "Caries Incidence (Decayed)\nN (%)", "Total Evaluated", "Primary Exact Test", "Secondary Asymptotic"]
    for j, h in enumerate(t5_headers):
        t5.rows[0].cells[j].paragraphs[0].text = h
    t5_data = [
        ["Score 0 (Completely Retained)", "85 (100.0%)", "0 (0.0%)", "85 (100.0%)", "Fisher-Freeman-Halton", "Pearson χ² = 118.00"],
        ["Score 1 (Partially Retained)", "28 (100.0%)", "0 (0.0%)", "28 (100.0%)", "Exact Test:", "df = 2"],
        ["Score 2 (Completely Missing)", "0 (0.0%)", "5 (100.0%)", "5 (100.0%)", "p < 0.0001 ***", "p < 0.0001"],
        ["Total Evaluated", "113 (95.8%)", "5 (4.2%)", "118 (100.0%)", "(Exact p = 5.75 × 10⁻⁹)", ""]
    ]
    for i, row_data in enumerate(t5_data):
        for j, val in enumerate(row_data):
            t5.rows[i+1].cells[j].paragraphs[0].text = val
    style_academic_table(t5, [2.0, 1.4, 1.4, 1.0, 1.3, 1.2], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    
    p_t5_note = doc.add_paragraph("Note: *** Highly statistically significant. Fisher-Freeman-Halton exact test was utilized as the primary test due to sparse cell frequencies (< 5) and structural zeros. Pearson Chi-Square is reported as supporting secondary analysis.")
    p_t5_note.runs[0].font.size = Pt(8.5)
    p_t5_note.runs[0].font.italic = True
    p_t5_note.paragraph_format.space_after = Pt(8)

    # FINDING 5: Baseline DMFT Stratification
    add_custom_heading("3.5 Finding 5: Baseline DMFT/DMFS Indices Stratified by Retention Outcome", level=2)
    p = doc.add_paragraph("Statistical Result: Due to the count/ordinal nature of caries indices and the small sample size in the complete loss group (Score 2, n = 5), the non-parametric Kruskal-Wallis H-test was conducted to compare baseline DMFT and DMFS across retention groups. No statistically significant differences were observed across the three retention outcomes (Baseline DMFT: H = 0.566, df = 2, p = 0.753; Baseline DMFS: H = 0.601, df = 2, p = 0.740). Supplementary one-way ANOVA yielded identical non-significant conclusions (DMFT: F = 0.187, p = 0.830; DMFS: F = 0.345, p = 0.709).")
    
    p = doc.add_paragraph("Clinical Deduction: This confirms that a patient's pre-existing baseline caries severity did not confound the technical retention of the sealant. Retention failure was governed by anatomical morphology and moisture isolation rather than individual baseline caries susceptibility.")

    # Table 6: Baseline DMFT Stratification
    add_custom_heading("Table 6. Baseline DMFT and DMFS indices across 3-Month Retention Groups", level=3)
    t6 = doc.add_table(rows=4, cols=6)
    t6_headers = ["Retention Group (Simonsen's)", "Sample Size (Teeth)", "Baseline DMFT (Mean ± SD)", "Baseline DMFS (Mean ± SD)", "Kruskal-Wallis Test (Primary)", "ANOVA (Supporting)"]
    for j, h in enumerate(t6_headers):
        t6.rows[0].cells[j].paragraphs[0].text = h
    t6_data = [
        ["Score 0 (Complete Retention)", "85", "2.24 ± 1.19", "4.52 ± 2.33", "DMFT: H = 0.566, p = 0.753", "DMFT: F = 0.187, p = 0.830"],
        ["Score 1 (Partial Retention)", "28", "2.39 ± 1.20", "4.89 ± 2.20", "DMFS: H = 0.601, p = 0.740", "DMFS: F = 0.345, p = 0.709"],
        ["Score 2 (Complete Loss)", "5", "2.20 ± 1.64", "4.20 ± 3.19", "(df = 2, Non-Significant)", "(df = 2, 115, NS)"]
    ]
    for i, row_data in enumerate(t6_data):
        for j, val in enumerate(row_data):
            t6.rows[i+1].cells[j].paragraphs[0].text = val
    style_academic_table(t6, [1.8, 0.7, 1.4, 1.4, 1.3, 1.3], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    
    p_t6_note = doc.add_paragraph("Note: Non-parametric Kruskal-Wallis H-test is reported as primary due to skewed distribution and small subgroup size in Score 2 (n = 5). Differences were non-significant (p > 0.05), confirming absence of baseline confounding.")
    p_t6_note.runs[0].font.size = Pt(8.5)
    p_t6_note.runs[0].font.italic = True
    p_t6_note.paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 3.6 SENSITIVITY ANALYSIS: CLUSTERING WITHIN PATIENTS (GEE)
    # -------------------------------------------------------------
    add_custom_heading("3.6 Methodological Sensitivity Analysis: Accounting for Clustered Teeth within Patients", level=2)
    p = doc.add_paragraph("Methodological Context: In dental clinical trials, evaluating multiple teeth per subject (nested/clustered design) can violate the assumption of observation independence if standard unadjusted tooth-level tests are used alone. In this cohort, 118 teeth were evaluated across 40 patients (cluster size range: 1–8 teeth, mean: 2.95 teeth/patient). To verify whether clustering within the child influenced statistical inferences, a sensitivity analysis was performed using Generalized Estimating Equations (GEE) with an exchangeable correlation structure and robust Huber-White standard errors.")
    
    p = doc.add_paragraph("Sensitivity Findings: The GEE modeling demonstrated that accounting for intra-subject correlation fully corroborated all primary findings:")
    
    bp_g1 = doc.add_paragraph(style='List Bullet')
    bp_g1.add_run("Tooth Type: ").bold = True
    bp_g1.add_run("Molars remained significantly associated with reduced complete retention compared to premolars (GEE robust z = -2.828, p = 0.005; Adjusted Odds Ratio = 0.322, 95% CI: 0.147–0.706), confirming the unadjusted Mann-Whitney U test (p = 0.002).")

    bp_g2 = doc.add_paragraph(style='List Bullet')
    bp_g2.add_run("Age Subgroup: ").bold = True
    bp_g2.add_run("Older children (> 10 years) maintained significantly higher complete retention compared to younger children (GEE robust z = 3.626, p = 0.0003; Adjusted Odds Ratio = 3.126, 95% CI: 1.688–5.790), confirming the unadjusted test (p = 0.006).")

    bp_g3 = doc.add_paragraph(style='List Bullet')
    bp_g3.add_run("Jaw Location: ").bold = True
    bp_g3.add_run("Upper vs. lower arch showed no significant effect under clustering (GEE robust z = -0.258, p = 0.797), fully agreeing with the unadjusted finding (p = 0.648).")

    # Table 7: GEE Sensitivity Model
    add_custom_heading("Table 7. Clustered sensitivity analysis using Generalized Estimating Equations (GEE) adjusting for 40 patient clusters (N = 118 teeth)", level=3)
    t7 = doc.add_table(rows=4, cols=6)
    t7_headers = ["Independent Predictor", "Unadjusted Tooth-Level p-value", "GEE Robust Coef (SE)", "Robust z-statistic", "Cluster-Adjusted p-value", "Adjusted Odds Ratio (95% CI)"]
    for j, h in enumerate(t7_headers):
        t7.rows[0].cells[j].paragraphs[0].text = h
    t7_data = [
        ["Tooth Type (Molars vs. Premolars)", "p = 0.002 **", "-1.133 (0.401)", "-2.828", "p = 0.005 **", "0.322 (0.147 – 0.706)"],
        ["Age Group (> 10 vs. ≤ 10 Years)", "p = 0.006 **", "+1.140 (0.314)", "+3.626", "p = 0.0003 ***", "3.126 (1.688 – 5.790)"],
        ["Jaw Location (Upper vs. Lower Arch)", "p = 0.648 (NS)", "-0.131 (0.508)", "-0.258", "p = 0.797 (NS)", "0.877 (0.324 – 2.373)"]
    ]
    for i, row_data in enumerate(t7_data):
        for j, val in enumerate(row_data):
            t7.rows[i+1].cells[j].paragraphs[0].text = val
    style_academic_table(t7, [1.8, 1.2, 1.2, 1.0, 1.1, 1.5], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
    
    p_t7_note = doc.add_paragraph("Note: GEE models specified Binomial family, logit link, exchangeable correlation structure across 40 patient clusters (mean cluster size = 2.95 teeth). All statistical conclusions remain completely robust.")
    p_t7_note.runs[0].font.size = Pt(8.5)
    p_t7_note.runs[0].font.italic = True
    p_t7_note.paragraph_format.space_after = Pt(10)

    # -------------------------------------------------------------
    # 4. KEY SYNTHESIS POINTS FOR MANUSCRIPT DISCUSSION
    # -------------------------------------------------------------
    add_custom_heading("4. Key Synthesis Points for the Discussion Section of the Manuscript", level=1)
    
    p = doc.add_paragraph("The following synthesis points provide refined academic text directly adaptable into the Discussion section of the manuscript:")
    
    bp_s1 = doc.add_paragraph(style='List Bullet')
    bp_s1.add_run("1. Satisfactory Short-Term Clinical Retention: ").bold = True
    bp_s1.add_run("Complete retention was observed in 72.0% of teeth, and satisfactory retention (Score 0 + Score 1) reached 95.8% at 3 months, reflecting high initial clinical efficacy of light-cured resin sealants in this pediatric cohort.")

    bp_s2 = doc.add_paragraph(style='List Bullet')
    bp_s2.add_run("2. Morphological Advantage of Premolars: ").bold = True
    bp_s2.add_run("Premolars showed significantly superior complete retention compared to molars (80.5% vs. 56.1%, p = 0.002; GEE cluster-adjusted p = 0.005). The shallower, wider groove morphology of premolars permits superior acid-etch conditioning and uniform resin flow compared to deep, constricted molar fissures.")

    bp_s3 = doc.add_paragraph(style='List Bullet')
    bp_s3.add_run("3. Enhanced Vigilance in Molar Application: ").bold = True
    bp_s3.add_run("Because all 5 complete sealant failures occurred exclusively in permanent molars (12.2% loss rate), clinicians must prioritize strict moisture control and consider dental isolation adjuncts, particularly when operating on partially erupted or distal molar surfaces.")

    bp_s4 = doc.add_paragraph(style='List Bullet')
    bp_s4.add_run("4. Age-Dependent Compliance & Crown Emergence: ").bold = True
    bp_s4.add_run("Children older than 10 years achieved significantly higher retention than younger children (81.5% vs. 60.4%, p = 0.006; GEE cluster-adjusted p = 0.0003). This is attributed to enhanced patient cooperation and complete clinical crown emergence free of overlying gingival tissue.")

    bp_s5 = doc.add_paragraph(style='List Bullet')
    bp_s5.add_run("5. Retention Status and Caries Association: ").bold = True
    bp_s5.add_run("Sealant retention was significantly associated with 3-month caries incidence (Fisher-Freeman-Halton exact test p < 0.0001). Intact and partially retained sealants provided a continuous physical barrier (100% caries-free), whereas complete dislodgement was associated with all newly observed carious lesions.")

    # 5. STRENGTHS AND METHODOLOGICAL LIMITATIONS
    add_custom_heading("5. Strengths and Methodological Limitations of the Study", level=1)
    
    add_custom_heading("5.1 Strengths", level=2)
    bp_st1 = doc.add_paragraph(style='List Bullet')
    bp_st1.add_run("Standardized Clinical Scoring: ").bold = True
    bp_st1.add_run("Application of Simonsen's internationally recognized criteria with rigorous baseline DMFT/DMFS recording.")
    bp_st2 = doc.add_paragraph(style='List Bullet')
    bp_st2.add_run("Methodological Rigor: ").bold = True
    bp_st2.add_run("Utilization of exact non-parametric tests (Fisher-Freeman-Halton, Kruskal-Wallis) combined with clustered sensitivity modeling (GEE) to verify statistical robustness.")

    add_custom_heading("5.2 Methodological Considerations & Limitations", level=2)
    bp_lim1 = doc.add_paragraph(style='List Bullet')
    bp_lim1.add_run("Tooth-Level Clustering: ").bold = True
    bp_lim1.add_run("Primary analyses were performed at the tooth level (N = 118) in consistency with the baseline study design. Because multiple teeth were treated per participant (mean: 2.95 teeth/child), potential intra-subject correlation was formally tested using GEE sensitivity analysis, which confirmed all primary conclusions.")
    
    bp_lim2 = doc.add_paragraph(style='List Bullet')
    bp_lim2.add_run("Follow-Up Duration and Confounders: ").bold = True
    bp_lim2.add_run("The 3-month evaluation window reflects early technical adhesion and short-term clinical outcomes. Longer follow-up (6, 12, and 24 months) and assessment of dietary and salivary covariates are recommended to evaluate long-term preventive performance.")
    
    bp_lim3 = doc.add_paragraph(style='List Bullet')
    bp_lim3.add_run("Absence of Untreated Control Group: ").bold = True
    bp_lim3.add_run("Due to ethical considerations regarding withholding standard preventive care in caries-active children, an unsealed control group was not included; thus, the relationship between sealant retention and caries incidence is interpreted as a strong clinical association.")

    # Save documents
    out_path1 = 'C:/Users/w/Desktop/احصاء بحث تخرج/Scopus_Statistical_Report_Final.docx'
    doc.save(out_path1)
    try:
        out_path2 = 'C:/Users/w/Desktop/احصاء بحث تخرج/Statistical_Findings_and_Clinical_Interpretations_Report.docx'
        doc.save(out_path2)
    except PermissionError:
        pass
    print('Doctor report successfully created.')

if __name__ == '__main__':
    create_doctor_report()
